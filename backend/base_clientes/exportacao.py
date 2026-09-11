"""
Gera o .xlsx de duas abas, no mesmo formato do tratamento manual antigo.

Duas conversões acontecem aqui e as duas têm motivo:

1. `item_contratado`, `cpf` e `proposta` saem como **texto**, com
   `number_format = "@"`. São identificadores, não números. Sem isso o Excel
   reintroduz a notação científica que esta feature existe para evitar.

2. Todo o resto faz o caminho inverso. O Supabase devolve `NUMERIC` como string
   (`"456902.55"`) e `DATE` como texto (`"2023-11-20"`); escrever isso direto
   produziria uma planilha em que a coluna de capital segurado não soma e a de
   vigência não ordena — e somar e ordenar é exatamente o que o usuário abre o
   arquivo para fazer.
"""

import io
from datetime import date, datetime
from typing import Any, Dict, List

import openpyxl

ABA_CLIENTES = [
    ("cpf", "CPF"), ("nome", "Nome do Cliente"), ("telefone", "Telefone / WhatsApp"),
    ("email", "E-mail"), ("total_capital_segurado", "Total Capital Segurado (R$)"),
    ("total_premio_mensal", "Total Prêmio Mensal (R$)"), ("qtd_propostas", "Qtd Propostas"),
    ("qtd_coberturas", "Qtd Coberturas"), ("profissao", "Profissão"),
    ("renda", "Renda Declarada (R$)"), ("data_nascimento", "Data Nascimento"),
    ("sexo", "Sexo"), ("qtd_filhos", "Qtd Filhos"), ("endereco", "Endereço Completo"),
]

ABA_COBERTURAS = [
    ("cpf", "CPF Cliente"), ("nome_segurado", "Nome Segurado"), ("proposta", "Nº Proposta"),
    ("item_contratado", "ID Item Cobertura"), ("produto", "Produto / Cobertura"),
    ("status_cobertura", "Status Cobertura"), ("capital_segurado", "Capital Segurado (R$)"),
    ("premio_atual", "Prêmio Atual (R$)"), ("premio_mensalizado", "Prêmio Mensalizado (R$)"),
    ("premio_anualizado", "Prêmio Anualizado (R$)"),
    ("premio_implantado_liquido", "Prêmio Implantado Líquido (R$)"),
    ("iof_implantado", "IOF Implantado (R$)"), ("periodicidade", "Periodicidade"),
    ("forma_pagamento", "Forma de Pagamento"), ("dia_vencimento", "Dia Vencimento"),
    ("data_entrada", "Data Entrada Proposta"), ("inicio_vigencia", "Início Vigência"),
    ("fim_vigencia", "Fim Vigência"), ("tempo_contribuicao", "Tempo Contribuição (Anos)"),
    ("prazo_decrescimo", "Prazo Decréscimo"), ("prazo_diferimento", "Prazo Diferimento"),
    ("am", "AM (Account Manager)"), ("produtor_principal", "Produtor Principal (Emissão)"),
    ("corretor_pj", "Corretor PJ"), ("corretor_estruturado", "Corretor Estruturado"),
    ("unidade_producao", "Unidade Produção"), ("oferta_comercial", "Oferta Comercial"),
]

# Identificadores: saem como texto e nunca são convertidos para número.
CAMPOS_IDENTIFICADOR = ("item_contratado", "cpf", "proposta")

FORMATO_MOEDA = "#,##0.00"
FORMATO_DATA = "DD/MM/YYYY"

CAMPOS_MOEDA = {
    "total_capital_segurado", "total_premio_mensal", "renda", "capital_segurado",
    "premio_atual", "premio_mensalizado", "premio_anualizado",
    "premio_implantado_liquido", "iof_implantado",
}
CAMPOS_DATA = {"data_nascimento", "data_entrada", "inicio_vigencia", "fim_vigencia"}


def _celula(campo: str, v: Any) -> Any:
    """Devolve o valor no tipo que o Excel precisa para somar e ordenar."""
    if v is None or v == "":
        return None
    if campo in CAMPOS_IDENTIFICADOR:
        return str(v)
    if campo in CAMPOS_MOEDA and isinstance(v, str):
        try:
            return float(v)
        except ValueError:
            return v
    if campo in CAMPOS_DATA and isinstance(v, str):
        try:
            return datetime.fromisoformat(v.replace("Z", "+00:00")).replace(tzinfo=None)
        except ValueError:
            return v
    # openpyxl não aceita datetime com fuso; o export é para leitura humana.
    if isinstance(v, datetime) and v.tzinfo is not None:
        return v.replace(tzinfo=None)
    return v


def _escrever(ws, colunas, linhas):
    ws.append([rotulo for _, rotulo in colunas])
    for item in linhas:
        ws.append([_celula(campo, item.get(campo)) for campo, _ in colunas])

    for col, (campo, _) in enumerate(colunas, start=1):
        if campo in CAMPOS_IDENTIFICADOR:
            formato = "@"
        elif campo in CAMPOS_MOEDA:
            formato = FORMATO_MOEDA
        elif campo in CAMPOS_DATA:
            formato = FORMATO_DATA
        else:
            continue
        for linha in range(2, ws.max_row + 1):
            ws.cell(row=linha, column=col).number_format = formato


def gerar_xlsx(clientes: List[Dict[str, Any]], coberturas: List[Dict[str, Any]]) -> bytes:
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "1_CLIENTES"
    # Desempata pelo CPF: nomes homônimos existem na carteira, e sem chave
    # estável a ordem muda a cada export sem que nada tenha mudado.
    _escrever(ws1, ABA_CLIENTES,
              sorted(clientes, key=lambda c: (c.get("nome") or "", c.get("cpf") or "")))
    ws2 = wb.create_sheet("2_APÓLICES_E_COBERTURAS")
    _escrever(ws2, ABA_COBERTURAS,
              sorted(coberturas,
                     key=lambda c: (c.get("nome_segurado") or "", c.get("item_contratado") or "")))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
